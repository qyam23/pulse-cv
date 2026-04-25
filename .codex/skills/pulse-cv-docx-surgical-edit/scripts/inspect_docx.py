import json
import sys
from collections import Counter
from pathlib import Path

from docx import Document


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        yield from iter_table_paragraphs(table)
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
        for table in section.header.tables:
            yield from iter_table_paragraphs(table)
        for table in section.footer.tables:
            yield from iter_table_paragraphs(table)


def inspect(path: Path):
    document = Document(str(path))
    paragraphs = list(iter_paragraphs(document))
    nonempty = [p for p in paragraphs if p.text.strip()]
    styles = Counter(p.style.name for p in paragraphs)
    return {
        "file": str(path),
        "exists": path.exists(),
        "bytes": path.stat().st_size,
        "paragraphs_total": len(paragraphs),
        "paragraphs_nonempty": len(nonempty),
        "tables_body": len(document.tables),
        "sections": len(document.sections),
        "styles": dict(styles.most_common()),
        "sample": [p.text.strip()[:160] for p in nonempty[:8]],
    }


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 2:
        raise SystemExit("Usage: inspect_docx.py <file.docx>")
    print(json.dumps(inspect(Path(sys.argv[1])), ensure_ascii=False, indent=2))
