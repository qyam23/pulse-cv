import copy
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph as PdfParagraph, SimpleDocTemplate, Spacer

BANNED_PHRASE_PATTERNS = [
    re.compile(r"\bif accurate\b", re.I),
    re.compile(r"\badd a bullet\b", re.I),
    re.compile(r"\bsuggested\b", re.I),
    re.compile(r"\brecommendation\b", re.I),
    re.compile(r"\bevidence-backed\b", re.I),
    re.compile(r"\bconsider adding\b", re.I),
    re.compile(r"\bif relevant\b", re.I),
    re.compile(r"\bif applicable\b", re.I),
    re.compile(r"\bimprove proof of\b", re.I),
    re.compile(r"\btailor this\b", re.I),
    re.compile(r"\bthe cv should\b", re.I),
    re.compile(r"\bthe candidate should\b", re.I),
    re.compile(r"\bedit plan\b", re.I),
    re.compile(r"\bplanner\b", re.I),
    re.compile(r"\bcoaching\b", re.I),
    re.compile(r"\banalysis\b", re.I),
    re.compile(r"\bapply recommendations?\b", re.I),
    re.compile(r"\brecruiter-ready\b", re.I),
]

BANNED_LINE_PATTERNS = [
    re.compile(r"^if\b", re.I),
    re.compile(r"^suggested\b", re.I),
    re.compile(r"^recommendation\b", re.I),
    re.compile(r"^consider\b", re.I),
    re.compile(r"^tailor\b", re.I),
    re.compile(r"^the cv should\b", re.I),
    re.compile(r"^the candidate should\b", re.I),
]

UNICODE_FONT_NAME = "PulseUnicode"
BOLD_UNICODE_FONT_NAME = "PulseUnicodeBold"


def register_unicode_fonts() -> Tuple[str, str]:
    regular_candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    bold_candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf"),
    ]

    regular = next((candidate for candidate in regular_candidates if candidate.exists()), None)
    bold = next((candidate for candidate in bold_candidates if candidate.exists()), None)
    if not regular:
        return "Helvetica", "Helvetica-Bold"

    if UNICODE_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(UNICODE_FONT_NAME, str(regular)))
    if bold and BOLD_UNICODE_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(BOLD_UNICODE_FONT_NAME, str(bold)))
    return UNICODE_FONT_NAME, BOLD_UNICODE_FONT_NAME if bold else UNICODE_FONT_NAME


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip().lower()


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    style = paragraph.style
    first_run_style = None
    if paragraph.runs:
      first = paragraph.runs[0]
      first_run_style = {
          "bold": first.bold,
          "italic": first.italic,
          "underline": first.underline,
          "font_name": first.font.name,
          "font_size": first.font.size,
      }

    p = paragraph._element
    for child in list(p):
        p.remove(child)

    run = paragraph.add_run(text)
    paragraph.style = style
    if first_run_style:
        run.bold = first_run_style["bold"]
        run.italic = first_run_style["italic"]
        run.underline = first_run_style["underline"]
        if first_run_style["font_name"]:
            run.font.name = first_run_style["font_name"]
        if first_run_style["font_size"]:
            run.font.size = first_run_style["font_size"]


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    try:
        new_para.style = paragraph.style
    except Exception:
        pass
    run = new_para.add_run(text)
    if paragraph.runs:
        ref = paragraph.runs[0]
        run.bold = ref.bold
        run.italic = ref.italic
        run.underline = ref.underline
        if ref.font.name:
            run.font.name = ref.font.name
        if ref.font.size:
            run.font.size = ref.font.size
    return new_para


def paragraph_contains(paragraph: Paragraph, target: str) -> bool:
    return normalize_text(target) in normalize_text(paragraph.text)


def find_paragraph_by_text(document: Document, target: str) -> Optional[Paragraph]:
    for paragraph in document.paragraphs:
        if paragraph_contains(paragraph, target):
            return paragraph
    return None


def all_document_text(document: Document) -> str:
    return "\n".join(p.text for p in document.paragraphs if p.text)


def extract_contact_lines(text: str) -> List[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    contacts = []
    for line in lines[:8]:
        if "@" in line or re.search(r"\+?\d[\d\-\s]{7,}", line):
            contacts.append(line)
    return contacts


def sanitize_recruiter_text(value: str) -> Tuple[str, List[str], bool]:
    warnings: List[str] = []
    if not value or not value.strip():
        return "", warnings, True

    kept_lines: List[str] = []
    for raw_line in value.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        banned = any(pattern.search(line) for pattern in BANNED_LINE_PATTERNS) or any(pattern.search(line) for pattern in BANNED_PHRASE_PATTERNS)
        if banned:
            warnings.append(f"Removed internal instruction line: {line}")
            continue
        kept_lines.append(line)

    text = "\n".join(kept_lines).strip()
    rejected = not text or any(pattern.search(text) for pattern in BANNED_PHRASE_PATTERNS)
    if rejected:
        warnings.append("Rejected contaminated recruiter-facing text before writing output.")
    return text, warnings, rejected


def validate_final_cv_output(text: str) -> None:
    violations = [pattern.pattern for pattern in BANNED_PHRASE_PATTERNS if pattern.search(text or "")]
    if violations:
        raise ValueError(f"Final CV output contains banned internal instruction language: {', '.join(violations)}")


def build_redline(report_path: Path, instructions: List[Dict[str, Any]], before_after: List[Dict[str, str]], warnings: List[str]) -> None:
    lines = ["# Pulse CV Change Report", ""]
    if warnings:
        lines.append("## Warnings")
        lines.extend([f"- {warning}" for warning in warnings])
        lines.append("")
    lines.append("## Applied changes")
    lines.append("")
    for index, instruction in enumerate(instructions, start=1):
        delta = next((item for item in before_after if item["id"] == instruction["id"]), None)
        lines.append(f"### {index}. {instruction['sectionLabel']} — {instruction['action']}")
        lines.append(f"- Rationale: {instruction['rationale']}")
        if instruction.get("linkedRequirementIds"):
            lines.append(f"- Linked requirements: {', '.join(instruction['linkedRequirementIds'])}")
        if delta:
            lines.append("- Before:")
            lines.append("")
            lines.append(f"```text\n{delta['before']}\n```")
            lines.append("- After:")
            lines.append("")
            lines.append(f"```text\n{delta['after']}\n```")
        lines.append("")
    report_path.write_text("\n".join(lines), encoding="utf-8")


def apply_docx(request: Dict[str, Any]) -> Tuple[str, str, List[str]]:
    source_path = Path(request["sourcePath"])
    output_path = Path(request["outputPath"])
    document = Document(str(source_path))
    instructions = request["editPlan"]["instructions"]
    before_after: List[Dict[str, str]] = []
    warnings: List[str] = []

    for instruction in instructions:
        action = instruction["action"]
        target = instruction.get("targetText") or ""
        replacement = instruction.get("replacementText") or ""
        anchor = instruction.get("insertionAnchor") or target
        replacement, sanitize_warnings, replacement_rejected = sanitize_recruiter_text(replacement)
        warnings.extend(sanitize_warnings)
        if instruction.get("replacementText") and replacement_rejected:
            warnings.append(f"Skipped contaminated instruction in section {instruction['sectionLabel']}.")
            continue

        if action in {"replace_phrase", "rewrite_bullet"} and target:
            paragraph = find_paragraph_by_text(document, target)
            if paragraph is None:
                warnings.append(f"Could not locate target text for instruction in section {instruction['sectionLabel']}.")
                continue
            before_after.append({"id": instruction["id"], "before": paragraph.text, "after": replacement})
            set_paragraph_text(paragraph, replacement)
        elif action == "insert_bullet":
            anchor_paragraph = find_paragraph_by_text(document, anchor) or (document.paragraphs[-1] if document.paragraphs else None)
            if anchor_paragraph is None:
                warnings.append(f"Could not locate anchor for insertion in section {instruction['sectionLabel']}.")
                continue
            inserted = insert_paragraph_after(anchor_paragraph, replacement)
            before_after.append({"id": instruction["id"], "before": "", "after": inserted.text})
        elif action == "tighten_heading":
            paragraph = find_paragraph_by_text(document, target)
            if paragraph is None:
                warnings.append(f"Could not locate heading to tighten in section {instruction['sectionLabel']}.")
                continue
            before_after.append({"id": instruction["id"], "before": paragraph.text, "after": replacement})
            set_paragraph_text(paragraph, replacement)

    original_text = request["resumeText"]
    updated_text = all_document_text(document)
    validate_final_cv_output(updated_text)
    for contact in extract_contact_lines(original_text):
        if normalize_text(contact) not in normalize_text(updated_text):
            warnings.append(f"Contact line may require manual verification: {contact}")

    document.save(str(output_path))
    build_redline(Path(request["redlinePath"]), instructions, before_after, warnings)
    return output_path.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", warnings


def build_pdf_story(text: str) -> List[Any]:
    styles = getSampleStyleSheet()
    body_font, heading_font = register_unicode_fonts()
    body = ParagraphStyle(
        "PulseBody",
        parent=styles["BodyText"],
        fontName=body_font,
        fontSize=10.5,
        leading=14,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "PulseHeading",
        parent=styles["Heading2"],
        fontName=heading_font,
        fontSize=12.5,
        leading=16,
        spaceBefore=10,
        spaceAfter=8,
    )
    story: List[Any] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 0.18 * cm))
            continue
        if len(line) < 40 and line == line.upper():
            story.append(PdfParagraph(line, heading))
            continue
        if re.match(r"^(summary|profile|experience|skills|education|projects|certifications|languages|תקציר|ניסיון|מיומנויות|השכלה)$", line, re.I):
            story.append(PdfParagraph(line, heading))
            continue
        escaped = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        if line.startswith(("-", "•", "*")):
            escaped = f"&bull; {escaped.lstrip('-•* ').strip()}"
        story.append(PdfParagraph(escaped, body))
    return story


def build_docx_from_text(output_path: Path, text: str) -> None:
    document = Document()
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph("")
            continue
        paragraph = document.add_paragraph(line)
        if len(line) < 45:
            for run in paragraph.runs:
                run.bold = True
    document.save(str(output_path))


def apply_pdf(request: Dict[str, Any]) -> Tuple[str, str, List[str]]:
    output_path = Path(request["outputPath"]).with_suffix(".docx")
    text = request["resumeText"]
    instructions = request["editPlan"]["instructions"]
    warnings = [
        "Source PDF was converted to an ATS-safe DOCX because Hebrew PDF regeneration can damage text extraction order."
    ]
    before_after: List[Dict[str, str]] = []
    updated_text = text

    for instruction in instructions:
        action = instruction["action"]
        target = instruction.get("targetText") or ""
        replacement = instruction.get("replacementText") or ""
        replacement, sanitize_warnings, replacement_rejected = sanitize_recruiter_text(replacement)
        warnings.extend(sanitize_warnings)
        if instruction.get("replacementText") and replacement_rejected:
            warnings.append(f"Skipped contaminated instruction in section {instruction['sectionLabel']}.")
            continue
        if action in {"replace_phrase", "rewrite_bullet"} and target and normalize_text(target) in normalize_text(updated_text):
            before_after.append({"id": instruction["id"], "before": target, "after": replacement})
            updated_text = re.sub(re.escape(target), replacement, updated_text, count=1)
        elif action == "insert_bullet":
            anchor = instruction.get("insertionAnchor") or ""
            if anchor and normalize_text(anchor) in normalize_text(updated_text):
                before_after.append({"id": instruction["id"], "before": "", "after": replacement})
                updated_text = updated_text.replace(anchor, f"{anchor}\n• {replacement}", 1)
            else:
                before_after.append({"id": instruction["id"], "before": "", "after": replacement})
                updated_text += f"\n• {replacement}"

    validate_final_cv_output(updated_text)
    build_docx_from_text(output_path, updated_text)
    build_redline(Path(request["redlinePath"]), instructions, before_after, warnings)
    return output_path.name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", warnings


def main() -> None:
    request_path = Path(sys.argv[1])
    request = json.loads(request_path.read_text(encoding="utf-8"))
    source_format = request["sourceFormat"]
    if source_format == "docx":
        output_file_name, mime_type, warnings = apply_docx(request)
    else:
        output_file_name, mime_type, warnings = apply_pdf(request)
    print(json.dumps({
        "outputFileName": output_file_name,
        "outputMimeType": mime_type,
        "outputPath": str(Path(request["outputPath"]).with_suffix(".docx")) if source_format == "pdf" else str(Path(request["outputPath"])),
        "warnings": warnings,
    }))


if __name__ == "__main__":
    main()
